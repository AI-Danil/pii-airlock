from __future__ import annotations

import multiprocessing
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from pii_airlock import documents
from pii_airlock.documents import extract_bytes, extract_path
from pii_airlock.models import AirlockError, ServiceBusyError


def test_txt_and_markdown() -> None:
    assert extract_bytes("Привет".encode(), ".txt") == "Привет"
    assert extract_bytes(b"# Heading", ".md") == "# Heading"


def test_docx_paragraphs_and_tables() -> None:
    doc = Document()
    doc.add_paragraph("Alice Carter")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "alice@example.test"
    buffer = BytesIO()
    doc.save(buffer)
    text = extract_bytes(buffer.getvalue(), ".docx")
    assert "Alice Carter" in text
    assert "alice@example.test" in text


def test_pdf_without_text_layer_is_rejected() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(AirlockError, match="OCR"):
        extract_bytes(buffer.getvalue(), ".pdf")


def test_pdf_with_text_layer_is_extracted() -> None:
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(72, 720, "Alice Carter")
    pdf.save()
    assert "Alice Carter" in extract_bytes(buffer.getvalue(), ".pdf")


def test_empty_damaged_and_oversized_inputs() -> None:
    with pytest.raises(AirlockError, match="empty"):
        extract_bytes(b"", ".txt")
    with pytest.raises(AirlockError, match="parser"):
        extract_bytes(b"not a zip", ".docx")
    with pytest.raises(AirlockError, match="20,000"):
        extract_bytes(("x" * 20_001).encode(), ".txt")


def test_docx_archive_expansion_is_bounded() -> None:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "content types")
        archive.writestr("word/document.xml", "x" * (25 * 1024 * 1024 + 1))
    with pytest.raises(AirlockError, match="expanded size"):
        extract_bytes(buffer.getvalue(), ".docx")


def test_pdf_page_count_is_bounded() -> None:
    writer = PdfWriter()
    for _ in range(101):
        writer.add_blank_page(width=200, height=200)
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(AirlockError, match="100-page"):
        extract_bytes(buffer.getvalue(), ".pdf")


def test_docx_archive_member_count_and_compression_ratio_are_bounded() -> None:
    many = BytesIO()
    with ZipFile(many, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("word/document.xml", "document")
        for index in range(documents.MAX_ARCHIVE_MEMBERS):
            archive.writestr(f"word/item-{index}.xml", "x")
    with pytest.raises(AirlockError, match="archive members"):
        extract_bytes(many.getvalue(), ".docx")

    compressed = BytesIO()
    with ZipFile(compressed, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("word/document.xml", "x" * (2 * 1024 * 1024))
    with pytest.raises(AirlockError, match="compression ratio"):
        extract_bytes(compressed.getvalue(), ".docx")


def test_binary_parser_runs_with_timeout_and_slot_limits(monkeypatch) -> None:
    doc = Document()
    doc.add_paragraph("Alice Carter")
    buffer = BytesIO()
    doc.save(buffer)

    assert documents._PARSER_SLOTS.acquire(blocking=False)
    assert documents._PARSER_SLOTS.acquire(blocking=False)
    try:
        with pytest.raises(ServiceBusyError, match="parser slots"):
            extract_bytes(buffer.getvalue(), ".docx")
    finally:
        documents._PARSER_SLOTS.release()
        documents._PARSER_SLOTS.release()

    monkeypatch.setattr(documents, "PARSER_TIMEOUT_SECONDS", 0)
    with pytest.raises(AirlockError, match="time limit"):
        extract_bytes(buffer.getvalue(), ".docx")
    assert not [child for child in multiprocessing.active_children() if child.name.startswith("SpawnProcess")]


def test_pdf_object_count_and_path_read_are_bounded(monkeypatch, tmp_path) -> None:
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(72, 720, "Alice Carter")
    pdf.save()
    monkeypatch.setattr(documents, "MAX_PDF_OBJECTS", 0)
    with pytest.raises(AirlockError, match="object safety"):
        documents._extract_binary_in_process(buffer.getvalue(), ".pdf")

    oversized = tmp_path / "oversized.txt"
    oversized.write_bytes(b"x" * (documents.MAX_BYTES + 1))
    with pytest.raises(AirlockError, match="5 MB"):
        extract_path(oversized)
