from __future__ import annotations

import multiprocessing
import os
import sys
import tempfile
import threading
from dataclasses import asdict, dataclass
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import BadZipFile, ZipFile

from .models import AirlockError, ServiceBusyError

MAX_BYTES = 5 * 1024 * 1024
MAX_CHARACTERS = 20_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
MAX_ARCHIVE_MEMBER_BYTES = 10 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 512
MAX_COMPRESSION_RATIO = 200
MAX_PDF_PAGES = 100
MAX_PDF_OBJECTS = 10_000
PARSER_TIMEOUT_SECONDS = 8.0
PARSER_MEMORY_BYTES = 512 * 1024 * 1024
PARSER_CPU_SECONDS = 6
SUPPORTED_SUFFIXES = {".txt", ".md", ".docx", ".pdf"}
_PARSER_SLOTS = threading.BoundedSemaphore(2)


@dataclass(frozen=True)
class DocumentManifest:
    format: str
    text_characters: int
    parser_isolation: str
    pages: int | None = None
    tables: int = 0
    headers: int = 0
    footers: int = 0
    unsupported_features: tuple[str, ...] = ()

    def public_dict(self) -> dict[str, object]:
        return asdict(self)


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    manifest: DocumentManifest


def extract_path(path: Path) -> str:
    return extract_path_with_manifest(path).text


def extract_path_with_manifest(path: Path) -> ExtractedDocument:
    if not path.exists() or not path.is_file():
        raise AirlockError(f"Document does not exist: {path}")
    with path.open("rb") as handle:
        data = handle.read(MAX_BYTES + 1)
    return extract_bytes_with_manifest(data, path.suffix)


def extract_bytes(data: bytes, suffix: str) -> str:
    return extract_bytes_with_manifest(data, suffix).text


def extract_bytes_with_manifest(data: bytes, suffix: str) -> ExtractedDocument:
    normalized_suffix = suffix.lower()
    if normalized_suffix not in SUPPORTED_SUFFIXES:
        raise AirlockError(f"Unsupported document type: {normalized_suffix or 'unknown'}")
    if not data:
        raise AirlockError("Document is empty.")
    if len(data) > MAX_BYTES:
        raise AirlockError("Document exceeds the 5 MB limit.")

    if normalized_suffix in {".txt", ".md"}:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise AirlockError("Document is damaged or cannot be decoded.") from exc
        manifest = DocumentManifest(
            format=normalized_suffix.removeprefix("."),
            text_characters=len(text),
            parser_isolation="not_required_plain_text",
        )
    else:
        extracted = _extract_isolated(data, normalized_suffix)
        text = extracted.text
        manifest = extracted.manifest

    cleaned = text.strip()
    if not cleaned:
        raise AirlockError("Document has no extractable text; OCR is outside V1.")
    if len(cleaned) > MAX_CHARACTERS:
        raise AirlockError("Extracted text exceeds the 20,000 character V1 limit.")
    if len(cleaned) != manifest.text_characters:
        manifest = DocumentManifest(**{**manifest.public_dict(), "text_characters": len(cleaned)})
    return ExtractedDocument(cleaned, manifest)


def _extract_binary_in_process(
    data: bytes,
    normalized_suffix: str,
    *,
    parser_isolation: str = "in_process_test_only",
) -> ExtractedDocument:
    try:
        if normalized_suffix == ".docx":
            from docx import Document

            _validate_docx_archive(data)
            document = Document(BytesIO(data))
            chunks = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    chunks.append("\t".join(cell.text for cell in row.cells))
            header_count = 0
            footer_count = 0
            seen_parts: set[str] = set()
            for section in document.sections:
                for label, part in (("header", section.header), ("footer", section.footer)):
                    part_name = str(part.part.partname)
                    if part_name in seen_parts:
                        continue
                    seen_parts.add(part_name)
                    part_chunks = [paragraph.text for paragraph in part.paragraphs]
                    for table in part.tables:
                        for row in table.rows:
                            part_chunks.append("\t".join(cell.text for cell in row.cells))
                    chunks.extend(part_chunks)
                    if any(chunk.strip() for chunk in part_chunks):
                        if label == "header":
                            header_count += 1
                        else:
                            footer_count += 1
            text = "\n".join(chunks)
            manifest = DocumentManifest(
                format="docx",
                text_characters=len(text),
                parser_isolation=parser_isolation,
                tables=len(document.tables),
                headers=header_count,
                footers=footer_count,
            )
        else:
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise AirlockError("Encrypted PDF documents are outside V1.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise AirlockError(f"PDF exceeds the {MAX_PDF_PAGES}-page V1 limit.")
            object_count = sum(len(objects) for objects in getattr(reader, "xref", {}).values())
            if object_count > MAX_PDF_OBJECTS:
                raise AirlockError(f"PDF exceeds the {MAX_PDF_OBJECTS}-object safety limit.")
            _validate_pdf_features(reader)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            manifest = DocumentManifest(
                format="pdf",
                text_characters=len(text),
                parser_isolation=parser_isolation,
                pages=len(reader.pages),
            )
    except AirlockError:
        raise
    except (UnicodeDecodeError, ValueError, KeyError, OSError) as exc:
        raise AirlockError("Document is damaged or cannot be decoded.") from exc
    except Exception as exc:
        raise AirlockError("Document parser rejected the file.") from exc
    return ExtractedDocument(text, manifest)


def _extract_isolated(data: bytes, normalized_suffix: str) -> ExtractedDocument:
    if not _PARSER_SLOTS.acquire(blocking=False):
        raise ServiceBusyError("All isolated document parser slots are busy; retry later.")
    context = multiprocessing.get_context("spawn")
    parent, child = context.Pipe(duplex=False)
    # The parent owns the scratch directory: a sandboxed worker may write inside
    # it but cannot remove it, so cleanup has to happen outside the boundary.
    temp_root = TemporaryDirectory(prefix="pii-airlock-parser-")
    process = context.Process(
        target=_parser_worker,
        args=(child, data, normalized_suffix, temp_root.name),
        daemon=True,
    )
    try:
        process.start()
        child.close()
        if not parent.poll(PARSER_TIMEOUT_SECONDS):
            process.terminate()
            process.join(timeout=2)
            raise AirlockError("Document parser exceeded its time limit.")
        status, payload = parent.recv()
        process.join(timeout=2)
        if status == "ok":
            return ExtractedDocument(
                text=str(payload["text"]),
                manifest=DocumentManifest(**payload["manifest"]),
            )
        message, code = payload
        if code == ServiceBusyError.code:
            raise ServiceBusyError(message)
        if os.getenv("PII_AIRLOCK_DEBUG_PARSER") == "1":
            raise AirlockError(f"{message} [{code}]")
        raise AirlockError(message)
    except EOFError as exc:
        raise AirlockError("Isolated document parser exited without a result.") from exc
    finally:
        parent.close()
        if process.is_alive():
            process.terminate()
            process.join(timeout=2)
        temp_root.cleanup()
        _PARSER_SLOTS.release()


def _parser_worker(connection, data: bytes, normalized_suffix: str, temp_root: str) -> None:
    try:
        work_dir = Path(temp_root) / "work"
        work_dir.mkdir(parents=True, exist_ok=True)
        _apply_resource_limits()
        isolation = _apply_parser_sandbox(work_dir, Path(temp_root))
        extracted = _extract_binary_in_process(data, normalized_suffix, parser_isolation=isolation)
        connection.send(
            (
                "ok",
                {"text": extracted.text, "manifest": extracted.manifest.public_dict()},
            )
        )
    except AirlockError as exc:
        connection.send(("error", (str(exc), getattr(exc, "code", "airlock_error"))))
    except BaseException as exc:
        connection.send(
            (
                "error",
                (
                    "Document parser failed inside its isolation boundary.",
                    f"parser_failed_{type(exc).__name__}:{str(exc)[:160]}",
                ),
            )
        )
    finally:
        connection.close()


def _apply_resource_limits() -> None:
    if os.name != "posix":
        return
    try:
        import resource

        resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
        resource.setrlimit(resource.RLIMIT_CPU, (PARSER_CPU_SECONDS, PARSER_CPU_SECONDS))
        if hasattr(resource, "RLIMIT_FSIZE"):
            resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))
        if hasattr(resource, "RLIMIT_NOFILE"):
            resource.setrlimit(resource.RLIMIT_NOFILE, (32, 32))
        if hasattr(resource, "RLIMIT_NPROC"):
            resource.setrlimit(resource.RLIMIT_NPROC, (0, 0))
        if hasattr(resource, "RLIMIT_AS"):
            resource.setrlimit(resource.RLIMIT_AS, (PARSER_MEMORY_BYTES, PARSER_MEMORY_BYTES))
    except (ImportError, OSError, ValueError):
        # The parent still enforces wall-clock termination and bounded input.
        return


def _apply_parser_sandbox(temp_dir: Path, writable_root: Path | None = None) -> str:
    """Remove ambient credentials and deny parser network/process activity.

    On macOS the worker also enters an OS sandbox. Other platforms retain the
    spawned-process and resource boundary and report that limitation explicitly.
    """
    writable_root = writable_root or temp_dir
    os.chdir(temp_dir)
    os.umask(0o077)
    os.environ.clear()
    os.environ.update({"HOME": str(temp_dir), "TMPDIR": str(temp_dir), "PATH": "/usr/bin:/bin"})
    tempfile.tempdir = str(temp_dir)

    def deny_ambient_capabilities(event: str, _args: tuple[object, ...]) -> None:
        if event.startswith("socket.") or event in {"subprocess.Popen", "os.system", "os.posix_spawn"}:
            raise PermissionError("Parser sandbox denied network or process execution.")

    sys.addaudithook(deny_ambient_capabilities)
    if sys.platform == "darwin" and _apply_macos_sandbox(writable_root):
        return "macos_sandbox+spawn+resource_limits+python_audit"
    _apply_linux_no_new_privs()
    return "spawn+resource_limits+python_audit;os_sandbox_unavailable"


def _apply_macos_sandbox(temp_dir: Path) -> bool:
    try:
        import ctypes

        library = ctypes.CDLL("/usr/lib/libsandbox.1.dylib")
        library.sandbox_init.argtypes = [ctypes.c_char_p, ctypes.c_uint64, ctypes.POINTER(ctypes.c_char_p)]
        library.sandbox_init.restype = ctypes.c_int
        error_buffer = ctypes.c_char_p()
        escaped_temp = str(temp_dir).replace('"', '\\"')
        profile = (
            "(version 1)(allow default)(deny network*)"
            "(deny process-exec)(deny file-write*)"
            f'(allow file-write* (subpath "{escaped_temp}"))'
        )
        result = library.sandbox_init(profile.encode("utf-8"), 0, ctypes.byref(error_buffer))
        if error_buffer.value and hasattr(library, "sandbox_free_error"):
            library.sandbox_free_error(error_buffer)
        return result == 0
    except (AttributeError, OSError, ValueError):
        return False


def _apply_linux_no_new_privs() -> None:
    if not sys.platform.startswith("linux"):
        return
    try:
        import ctypes

        libc = ctypes.CDLL(None)
        libc.prctl(38, 1, 0, 0, 0)
    except (AttributeError, OSError, ValueError):
        return


def _validate_docx_archive(data: bytes) -> None:
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise AirlockError(f"DOCX contains more than {MAX_ARCHIVE_MEMBERS} archive members.")
            names = {member.filename for member in members}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise AirlockError("DOCX package is missing required document parts.")
            total_uncompressed = sum(member.file_size for member in members)
            if total_uncompressed > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise AirlockError("DOCX expanded size exceeds the 25 MB safety limit.")
            for member in members:
                if member.flag_bits & 0x1:
                    raise AirlockError("Encrypted DOCX archive members are outside V1.")
                if member.file_size > MAX_ARCHIVE_MEMBER_BYTES:
                    raise AirlockError("A DOCX archive member exceeds the 10 MB safety limit.")
                ratio = member.file_size / max(1, member.compress_size)
                if member.file_size > 1024 * 1024 and ratio > MAX_COMPRESSION_RATIO:
                    raise AirlockError("DOCX compression ratio exceeds the safety limit.")
            unsupported: set[str] = set()
            for name in names:
                lowered = name.casefold()
                if lowered.startswith(("word/embeddings/", "word/activex/")):
                    unsupported.add("embedded_object")
                if lowered.startswith("word/comments"):
                    unsupported.add("comments")
                if lowered in {"word/footnotes.xml", "word/endnotes.xml"}:
                    unsupported.add("footnotes_or_endnotes")
                if lowered.endswith("vbaproject.bin"):
                    unsupported.add("macro")
            xml_markers = {
                b"<w:altChunk": "external_or_alternative_chunk",
                b"<w:txbxContent": "text_box",
                b"<w:vanish": "hidden_text",
            }
            for name in names:
                lowered = name.casefold()
                if not lowered.endswith((".xml", ".rels")):
                    continue
                content = archive.read(name)
                # A DOCX part never needs a DTD; refusing one keeps entity
                # expansion and external-entity fetches out of the parser.
                probe = content[:4096].lstrip()
                if b"<!DOCTYPE" in probe or b"<!ENTITY" in content:
                    raise AirlockError("DOCX declares an XML document type or entity, which is rejected.")
                if not lowered.startswith("word/") or not lowered.endswith(".xml"):
                    continue
                for marker, feature in xml_markers.items():
                    if marker in content:
                        unsupported.add(feature)
            if unsupported:
                features = ", ".join(sorted(unsupported))
                raise AirlockError(
                    f"DOCX contains unsupported content that would make extraction incomplete: {features}."
                )
    except BadZipFile as exc:
        raise AirlockError("Document parser rejected the file.") from exc


def _validate_pdf_features(reader) -> None:
    root = reader.trailer["/Root"].get_object()
    unsupported: set[str] = set()
    names = root.get("/Names")
    if names is not None:
        names = names.get_object()
        if names.get("/EmbeddedFiles") is not None:
            unsupported.add("embedded_files")
        if names.get("/JavaScript") is not None:
            unsupported.add("javascript")
    if root.get("/OpenAction") is not None or root.get("/AA") is not None:
        unsupported.add("automatic_action")
    if root.get("/OCProperties") is not None:
        unsupported.add("optional_content_layers")
    form = root.get("/AcroForm")
    if form is not None:
        form = form.get_object()
        unsupported.add("interactive_form")
        if form.get("/XFA") is not None:
            unsupported.add("xfa_form")
    if any(page.get("/Annots") is not None for page in reader.pages):
        unsupported.add("annotations")
    if unsupported:
        features = ", ".join(sorted(unsupported))
        raise AirlockError(f"PDF contains unsupported active or hidden content: {features}.")
